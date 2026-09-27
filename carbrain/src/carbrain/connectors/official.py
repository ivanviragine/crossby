"""Official statistics: Banco Central (SGS), IBGE (IPCA), ANP (fuel), ANFAVEA (registrations)."""

from __future__ import annotations

import calendar
import io
import json
import re
from collections.abc import Iterable, Iterator
from datetime import date, datetime
from typing import Any, ClassVar

import openpyxl

from carbrain.archive import Snapshot
from carbrain.connectors.base import Connector, ParseContext
from carbrain.geo import municipio_id, uf_code
from carbrain.http import Http
from carbrain.models import Observation
from carbrain.resolve import normalize


class ParseError(ValueError):
    pass


def month_period(year: int, month: int) -> tuple[date, date]:
    return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])


# --- Banco Central: vehicle loan rates -------------------------------------------------

SGS_SERIES: dict[int, tuple[str, str]] = {
    25471: ("vehicle_loan_rate_pf_monthly", "% a.m."),
    20749: ("vehicle_loan_rate_pf_annual", "% a.a."),
}


class BcbSgs(Connector):
    """Average interest rate on new vehicle loans to individuals (free resources)."""

    source_id = "bcb_sgs"
    parser_version = "1"
    base = "https://api.bcb.gov.br/dados/serie/bcdata.sgs.{code}/dados"

    def __init__(self, start: date = date(2011, 3, 1), today: date | None = None) -> None:
        self.start = start
        self.today = today or date.today()

    def discover(self, http: Http) -> list[str]:
        # `ultimos/N` is capped at 20 values, so ask for a date range. The end date is the
        # end of next year, which keeps the URL stable for a whole year.
        end = date(self.today.year + 1, 12, 31)
        return [
            f"{self.base.format(code=code)}?formato=json"
            f"&dataInicial={self.start:%d/%m/%Y}&dataFinal={end:%d/%m/%Y}"
            for code in SGS_SERIES
        ]

    def parse(self, snapshot: Snapshot, ctx: ParseContext) -> Iterator[Observation]:
        found = re.search(r"sgs\.(\d+)", snapshot.url)
        if found is None:
            raise ParseError(f"No SGS series code in {snapshot.url}")
        code = int(found.group(1))
        metric, unit = SGS_SERIES[code]
        payload = json.loads(snapshot.path.read_bytes())
        if isinstance(payload, dict):
            raise ParseError(f"SGS {code} returned an error: {payload.get('erro', payload)}")
        for item in payload:
            day = datetime.strptime(item["data"], "%d/%m/%Y").date()
            start, end = month_period(day.year, day.month)
            yield Observation(
                source_id=self.source_id,
                metric=metric,
                subject_type="national",
                subject_id="BR",
                dims={"sgs_series": code},
                as_of_start=start,
                as_of_end=end,
                value=float(item["valor"]),
                unit=unit,
            )


# --- IBGE: IPCA index ------------------------------------------------------------------


class IbgeIpca(Connector):
    """IPCA price index number (Dec 1993 = 100), the deflator for real prices."""

    source_id = "ibge_ipca"
    parser_version = "1"
    url = "https://apisidra.ibge.gov.br/values/t/1737/n1/all/v/2266/p/all"

    def discover(self, http: Http) -> list[str]:
        return [self.url]

    def parse(self, snapshot: Snapshot, ctx: ParseContext) -> Iterator[Observation]:
        rows = json.loads(snapshot.path.read_bytes())
        if not rows or rows[0].get("V") != "Valor":
            raise ParseError("Unexpected SIDRA layout: first row is not the header")
        for row in rows[1:]:
            value = row["V"]
            if not re.fullmatch(r"-?\d+(\.\d+)?", value):
                continue  # SIDRA uses '...', '-', 'X' for missing or confidential values
            period = row["D3C"]
            start, end = month_period(int(period[:4]), int(period[4:]))
            yield Observation(
                source_id=self.source_id,
                metric="ipca_index",
                subject_type="national",
                subject_id="BR",
                as_of_start=start,
                as_of_end=end,
                value=float(value),
                unit="index (Dec 1993 = 100)",
            )


# --- ANP: weekly retail fuel prices ----------------------------------------------------

ANP_METRICS = {
    "PRECO MEDIO REVENDA": "fuel_price_retail_avg",
    "PRECO MINIMO REVENDA": "fuel_price_retail_min",
    "PRECO MAXIMO REVENDA": "fuel_price_retail_max",
    "NUMERO DE POSTOS PESQUISADOS": "fuel_stations_surveyed",
}


class AnpWeekly(Connector):
    """Weekly survey summary: average, min and max pump price by product and place."""

    source_id = "anp_weekly"
    parser_version = "1"
    listing = (
        "https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/"
        "levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas"
    )
    # Sheet -> how to identify the place in each row.
    sheets: ClassVar[dict[str, str]] = {
        "BRASIL": "national",
        "REGIOES": "region",
        "ESTADOS": "uf",
        "MUNICIPIOS": "municipio",
    }

    def __init__(self, weeks: int = 1) -> None:
        self.weeks = weeks

    def discover(self, http: Http) -> list[str]:
        return self.links_from_listing(http.get(self.listing).text)[: self.weeks][::-1]

    @staticmethod
    def links_from_listing(html: str) -> list[str]:
        """Weekly summary links, newest first, as the page lists them.

        File names are inconsistent (e.g. `..._2026-08-30-2026-09-5.xlsx`), so links are
        read from the page instead of being built from dates.
        """
        seen: dict[str, None] = {}
        for href in re.findall(r'href="([^"]*resumo_semanal_lpc[^"]*\.xlsx)"', html):
            seen.setdefault(href, None)
        return list(seen)

    def parse(self, snapshot: Snapshot, ctx: ParseContext) -> Iterator[Observation]:
        wb = openpyxl.load_workbook(io.BytesIO(snapshot.path.read_bytes()), read_only=True)
        found = False
        for sheet, subject_type in self.sheets.items():
            if sheet not in wb.sheetnames:
                continue
            found = True
            yield from self._parse_sheet(wb[sheet].iter_rows(values_only=True), subject_type)
        if not found:
            raise ParseError(f"No known sheets in {snapshot.url}: {wb.sheetnames}")

    def _parse_sheet(
        self, rows: Iterable[tuple[Any, ...]], subject_type: str
    ) -> Iterator[Observation]:
        header: list[str] | None = None
        for row in rows:
            if header is None:
                if row and row[0] == "DATA INICIAL":
                    header = [normalize(str(c)) if c is not None else "" for c in row]
                continue
            if not row or row[0] is None:
                continue
            rec = dict(zip(header, row, strict=False))
            start = _as_date(rec["DATA INICIAL"])
            end = _as_date(rec["DATA FINAL"])
            product = normalize(str(rec["PRODUTO"]))
            subject_id = self._subject(rec, subject_type)
            for column, metric in ANP_METRICS.items():
                value = rec.get(column)
                if value is None or value == "-":
                    continue
                yield Observation(
                    source_id=self.source_id,
                    metric=metric,
                    subject_type=subject_type,
                    subject_id=subject_id,
                    dims={"product": product},
                    as_of_start=start,
                    as_of_end=end,
                    value=float(value),
                    unit="stations"
                    if metric == "fuel_stations_surveyed"
                    else rec["UNIDADE DE MEDIDA"],
                )
        if header is None:
            raise ParseError("ANP sheet has no 'DATA INICIAL' header row")

    @staticmethod
    def _subject(rec: dict[str, Any], subject_type: str) -> str:
        if subject_type == "national":
            return "BR"
        if subject_type == "region":
            return normalize(str(rec["REGIAO"]))
        if subject_type == "uf":
            return uf_code(str(rec["ESTADOS"]))
        return municipio_id(str(rec["ESTADO"]), str(rec["MUNICIPIO"]))


def _as_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()


# --- ANFAVEA: monthly registrations by segment -----------------------------------------

MONTHS_PT = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]


class Anfavea(Connector):
    """Monthly registrations (emplacamentos) by segment: domestic, imported and total."""

    source_id = "anfavea"
    parser_version = "1"
    missing_ok = True
    url = "https://anfavea.com.br/docs/siteautoveiculos{year}.xlsx"

    def __init__(self, years: list[int] | None = None) -> None:
        # Last year too, so year-on-year comparisons work from the first sync.
        this_year = date.today().year
        self.years = years or [this_year - 1, this_year]

    def discover(self, http: Http) -> list[str]:
        # The Excel page renders links with JavaScript, so the per-year URL is used.
        return [self.url.format(year=y) for y in sorted(self.years)]

    def parse(self, snapshot: Snapshot, ctx: ParseContext) -> Iterator[Observation]:
        wb = openpyxl.load_workbook(io.BytesIO(snapshot.path.read_bytes()), read_only=True)
        sheet = next((s for s in wb.sheetnames if normalize(s) == "I. EMPLACAMENTO"), None)
        if sheet is None:
            raise ParseError(f"No 'I. Emplacamento' sheet in {snapshot.url}: {wb.sheetnames}")
        yield from self.parse_rows(wb[sheet].iter_rows(values_only=True))

    def parse_rows(self, rows: Iterable[tuple[Any, ...]]) -> Iterator[Observation]:
        origin: str | None = None
        year: int | None = None
        month_cols: list[int] = []
        published: set[int] | None = None
        parent = ""
        for row in rows:
            cells = list(row) + [None] * 4
            title = cells[1]
            if isinstance(title, str) and normalize(title).startswith("EMPLACAMENTO"):
                origin = _origin(title)
                year, month_cols, parent = None, [], ""
                continue
            if origin is None:
                continue
            if cells[2] == "Unidades" and isinstance(cells[3], int):
                year = cells[3]
                continue
            if cells[3] == "Jan":
                month_cols = [i for i, c in enumerate(cells) if c in MONTHS_PT]
                published = None
                continue
            if year is None or not month_cols:
                continue
            level1, level2 = cells[1], cells[2]
            if isinstance(level1, str) and normalize(level1).startswith("FONTE"):
                origin = None
                continue
            if isinstance(level1, str):
                parent = _slug(level1)
                segment = parent
            elif isinstance(level2, str):
                segment = f"{parent}/{_slug(level2)}"
            else:
                continue
            if published is None:
                # The first row of a block is its total. Months not published yet are
                # empty in some files and 0 in others (2026), so a month counts as
                # published only when the block total is positive. A real 0 in a
                # sub-row (e.g. imported buses) is kept.
                published = {m for m, col in enumerate(month_cols, 1) if _positive(cells[col])}
            for month, col in enumerate(month_cols, start=1):
                value = cells[col]
                if month not in published or not isinstance(value, int | float):
                    continue
                start, end = month_period(year, month)
                yield Observation(
                    source_id=self.source_id,
                    metric="registrations",
                    subject_type="segment",
                    subject_id=segment,
                    dims={"origin": origin},
                    as_of_start=start,
                    as_of_end=end,
                    value=float(value),
                    unit="vehicles",
                )


def _positive(value: Any) -> bool:
    return isinstance(value, int | float) and value > 0


def _origin(title: str) -> str:
    norm = normalize(title)
    if "NACIONAIS" in norm:
        return "domestic"
    if "IMPORTADOS" in norm:
        return "imported"
    return "total"


def _slug(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", normalize(label).lower()).strip("_")
